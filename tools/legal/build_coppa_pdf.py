"""Build the served COPPA disclosure asset from the counsel-approved source.

Reads the dropped source docx, applies the two approved edits, writes the
edited docx of record + the served PDF, and prints the PDF's sha256 so the
exact hash can be pinned in src/lib/consent/text.ts.

Edits (the ONLY changes to the counsel text):
  1. §2 — remove "school name, " from the not-collected list (school is now
     collected; counsel cleared it).
  2. §10 — fill "[Insert Privacy Contact Email]" with privacy@samnewyork.com.

The PDF reproduces the disclosure notice (title through §10). The trailing
"Checkbox:" / "Button text:" lines in the source are UI authoring annotations,
not disclosure prose, and are intentionally not part of the served document;
the checkbox attestation itself lives in src/lib/consent/text.ts (CONSENT_TEXT).
"""

import datetime
import hashlib
import re
import sys

from docx import Document
from fpdf import FPDF

SRC = "docs/legal/COPPA Disclosure.docx"
EDITED_DOCX = "docs/legal/COPPA_Disclosure.docx"
PDF_OUT = "public/legal/coppa-disclosure-v1.pdf"

PRIVACY_EMAIL = "privacy@samnewyork.com"


def replace_in_paragraph(paragraph, old, new):
    """Run-aware text replace within one Normal paragraph (no rich runs)."""
    full = "".join(run.text for run in paragraph.runs)
    if old not in full:
        return False
    full = full.replace(old, new)
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = full
    else:
        paragraph.add_run(full)
    return True


def ascii_punct(s):
    return (
        s.replace("’", "'").replace("‘", "'")
        .replace("“", '"').replace("”", '"')
        .replace("–", "-").replace("—", "-")
        .replace(" ", " ")
    )


def main():
    doc = Document(SRC)

    edits = {"school": False, "email": False}
    for p in doc.paragraphs:
        if "We do not ask for the child" in p.text and "school name" in p.text:
            edits["school"] = replace_in_paragraph(p, "school name, ", "")
        if "[Insert Privacy Contact Email]" in p.text:
            edits["email"] = replace_in_paragraph(
                p, "[Insert Privacy Contact Email]", PRIVACY_EMAIL
            )

    if not edits["school"] or not edits["email"]:
        print(f"FATAL: edits not all applied: {edits}", file=sys.stderr)
        sys.exit(1)

    doc.save(EDITED_DOCX)

    # ---- Render the served PDF (title .. end of §10) ----
    pdf = FPDF(format="Letter", unit="mm")
    pdf.set_margins(left=22, top=22, right=22)
    pdf.set_auto_page_break(auto=True, margin=22)
    pdf.set_creation_date(datetime.datetime(2026, 6, 18, 0, 0, 0))
    pdf.set_title("COPPA Disclosure and Parental Consent")
    pdf.set_author("S.A.M New York")
    pdf.set_subject("coppa-disclosure-v1")
    pdf.set_creator("atlas-ai/tools/legal/build_coppa_pdf.py")
    pdf.add_page()

    section_re = re.compile(r"^\d+\.\s")
    for idx, p in enumerate(doc.paragraphs):
        text = ascii_punct(p.text).strip()
        if text.startswith("Checkbox:"):
            break  # authoring annotation — not part of the served disclosure
        if not text:
            continue
        if idx == 0:  # title
            pdf.set_font("Helvetica", "B", 16)
            pdf.multi_cell(0, 8, text, align="C")
            pdf.ln(3)
        elif section_re.match(text):
            pdf.ln(2)
            pdf.set_font("Helvetica", "B", 12)
            pdf.multi_cell(0, 6, text)
            pdf.ln(1)
        else:
            pdf.set_font("Helvetica", "", 11)
            pdf.multi_cell(0, 5.5, text)
            pdf.ln(1.5)

    pdf.output(PDF_OUT)

    digest = hashlib.sha256(open(PDF_OUT, "rb").read()).hexdigest()
    print(f"edited_docx={EDITED_DOCX}")
    print(f"pdf={PDF_OUT}")
    print(f"sha256={digest}")


if __name__ == "__main__":
    main()
